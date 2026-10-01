#!/usr/bin/env python3
"""
SL & OTIF core (openpyxl, pure Python).

Rule set (same as the sample "2026_<PRINCIPAL>_SL & OTIF.xlsx"):
  PERIODE          = month of PLANNED RECEIPT DATE
  Planned + N      = PLANNED RECEIPT DATE + N days        (tolerance, default 7)
  GAP              = (PLANNED RECEIPT DATE + N) - ARRIVAL DATE ; blank arrival -> -1
  Remarks          = GAP >= 0 -> "On Time" else "Late"
  SUMMARY per month: Late {PO,ARR} / On Time {PO,ARR} / Total / FULFIL=G/F / OTIF=E/F

The raw worksheet is kept AS-IS (headers + formatting). Helper columns are
APPENDED as live formulas (DATEVALUE-based, since dates are ISO text) and the
SUMMARY is formula-driven (SUMIFS). Cached values are injected so numbers show
immediately while formulas remain editable/recomputable.
"""
import datetime, io, re, shutil, zipfile
from copy import copy
from xml.sax.saxutils import escape

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

MONTHS = ["Jan","Feb","Mar","Apr","May","Jun","Jul","Aug","Sep","Oct","Nov","Dec"]
TOL = 7
NUMACC = '_(* #,##0_);_(* \\(#,##0\\);_(* "-"??_);_(@_)'
DATEFMT = 'yyyy\\-mm\\-dd;@'
THIN = Side(style="thin")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HFILL = PatternFill("solid", fgColor="FFF2CC")
EPOCH = datetime.date(1899, 12, 30)

FIELDS = {
    "pcode": ["principal code","principal _ code"], "pname": ["principal name","principal _ name"],
    "periode": ["periode","period"], "contract": ["contract"], "po": ["po number"], "note": ["po note"],
    "line": ["line no","line _ no"], "rel": ["rel no","rel _ no"], "podate": ["po date","po _ date"],
    "prod": ["product code","product _ code"], "desc": ["description"], "poqty": ["po qty"],
    "hna": ["hna um"], "povalue": ["po value"], "plan": ["planned receipt date","planned receipt _ date"],
    "status": ["po line status","po line _ status"], "lpb": ["lpb number"], "arrdate": ["arrival date"],
    "arrqty": ["arrival qty"], "arrvalue": ["arrival value"],
}


def norm(s): return re.sub(r"\s+", " ", str(s or "").strip().lower())


def to_date(v):
    if v is None: return None
    if isinstance(v, datetime.datetime): return v.date()
    if isinstance(v, datetime.date): return v
    s = str(v).strip()
    for fmt in ("%Y-%m-%d","%d/%m/%Y","%Y-%m-%d %H:%M:%S"):
        try: return datetime.datetime.strptime(s, fmt).date()
        except ValueError: pass
    return None


def num(v):
    if v is None: return 0
    if isinstance(v, (int, float)): return v
    try: return float(str(v).replace(",", "").strip())
    except ValueError: return 0


def serial(d): return (d - EPOCH).days


def _pick_sheet(wb):
    for n in wb.sheetnames:
        if norm(n) == "worksheet": return wb[n]
    return wb.worksheets[-1]


def _parse(wb, tol):
    ws = _pick_sheet(wb)
    hrow = 5
    for r in range(1, min(ws.max_row, 25) + 1):
        vals = [norm(ws.cell(r, c).value) for c in range(1, ws.max_column + 1)]
        if "po qty" in vals and any("principal" in v for v in vals):
            hrow = r; break
    cmap = {}
    for c in range(1, ws.max_column + 1):
        h = norm(ws.cell(hrow, c).value)
        for k, names in FIELDS.items():
            if k not in cmap and h in names: cmap[k] = c
    po_col = cmap.get("po", 5)
    last = hrow
    for r in range(hrow + 1, ws.max_row + 1):
        if ws.cell(r, po_col).value not in (None, ""): last = r
    rows = []
    for r in range(hrow + 1, last + 1):
        g = lambda k: ws.cell(r, cmap[k]).value if k in cmap else None
        plan = to_date(g("plan")); arr = to_date(g("arrdate"))
        gap = -1 if (not plan or not arr) else (plan + datetime.timedelta(days=tol) - arr).days
        rows.append(dict(row=r, plan=plan, arr=arr, gap=gap, on_time=(gap >= 0),
                         poqty=num(g("poqty")), arrqty=num(g("arrqty")), pname=g("pname")))
    return ws, hrow, last, cmap, rows


def _append_helpers(ws, hrow, rows, tol, planL, arrL, cache):
    start = ws.max_column + 1
    periodL, tolL, gapL, remL = (get_column_letter(start + i) for i in range(4))
    heads = ["PERIODE", f"Planned receipt + {tol} hari", "GAP", "Remarks"]
    hstyle = ws.cell(hrow, 1)
    for i, h in enumerate(heads):
        c = ws.cell(hrow, start + i, h)
        c._style = copy(hstyle._style)
        c.font = Font(bold=True); c.fill = HFILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
        ws.column_dimensions[get_column_letter(start + i)].width = 14 if i else 11
    for x in rows:
        r = x["row"]; left = ws.cell(r, start - 1)
        cells = [
            (start,     f'=IF({planL}{r}="","",TEXT(DATEVALUE({planL}{r}),"MMM"))',
             (MONTHS[x["plan"].month - 1] if x["plan"] else ""), True, None),
            (start + 1, f'=IF({planL}{r}="","",DATEVALUE({planL}{r})+{tol})',
             (serial(x["plan"]) + tol if x["plan"] else None), False, DATEFMT),
            (start + 2, f'=IF(OR({planL}{r}="",{arrL}{r}=0),-1,{tolL}{r}-DATEVALUE({arrL}{r}))',
             x["gap"], False, None),
            (start + 3, f'=IF({gapL}{r}>=0,"On Time","Late")', ("On Time" if x["on_time"] else "Late"), True, None),
        ]
        for col, formula, val, is_str, fmt in cells:
            c = ws.cell(r, col); c._style = copy(left._style); c.border = BORDER
            c.alignment = Alignment(horizontal="center", vertical="center")
            if fmt: c.number_format = fmt
            c.value = formula
            if val is not None:
                cache.setdefault(ws.title, {})[f"{get_column_letter(col)}{r}"] = (val, is_str)
    return periodL, tolL, gapL, remL


def _build_summary(wb, order, agg, grand, title, year, wsname, poL, arrL, periodL, remL, first, last, cache):
    sm = wb.create_sheet("SUMMARY", 0)
    for i, w in enumerate([13.1,16.3,19.3,14.1,20.3,19.1,27.7,9.1,9.1], 1):
        sm.column_dimensions[get_column_letter(i)].width = w
    sm["A1"] = title; sm["A1"].font = Font(bold=True)
    sm["A2"] = year
    labels = {"B3":"Column Labels","H3":"%","I3":"%","B4":"Late","D4":"On Time",
              "F4":"Total Sum of PO QTY","G4":"Total Sum of ARRIVAL QTY","H4":"FULFIL","I4":"OTIF",
              "A5":"Row Labels","B5":"Sum of PO QTY","C5":"Sum of ARRIVAL QTY","D5":"Sum of PO QTY",
              "E5":"Sum of ARRIVAL QTY","H5":"%","I5":"%"}
    for a, v in labels.items(): sm[a] = v
    for row in (3, 4, 5):
        for col in range(2, 10):
            cc = sm.cell(row, col); cc.border = BORDER
            if cc.value not in (None, "%"): cc.font = Font(bold=True)
            cc.alignment = Alignment(horizontal="center" if cc.value == "%" else "left")
    q = "'" + wsname.replace("'", "''") + "'!"
    rng = lambda L: f"{q}${L}${first}:${L}${last}"

    def sumifs(valL, mc, remark):
        if mc: return f'=SUMIFS({rng(valL)}, {rng(periodL)}, $A{mc}, {rng(remL)}, "{remark}")'
        return f'=SUMIFS({rng(valL)}, {rng(remL)}, "{remark}")'

    def write_row(r, mc, d):
        sc = cache.setdefault("SUMMARY", {})
        F = d["lateL"] + d["ontL"]; G = d["lateS"] + d["ontS"]
        sm.cell(r, 2, sumifs(poL, mc, "Late"));     sc[f"B{r}"] = (d["lateL"], False)
        sm.cell(r, 3, sumifs(arrL, mc, "Late"));    sc[f"C{r}"] = (d["lateS"], False)
        sm.cell(r, 4, sumifs(poL, mc, "On Time"));  sc[f"D{r}"] = (d["ontL"], False)
        sm.cell(r, 5, sumifs(arrL, mc, "On Time")); sc[f"E{r}"] = (d["ontS"], False)
        sm.cell(r, 6, f"=B{r}+D{r}"); sc[f"F{r}"] = (F, False)
        sm.cell(r, 7, f"=C{r}+E{r}"); sc[f"G{r}"] = (G, False)
        sm.cell(r, 8, f"=G{r}/F{r}"); sc[f"H{r}"] = ((G / F) if F else 0, False)
        sm.cell(r, 9, f"=E{r}/F{r}"); sc[f"I{r}"] = ((d["ontS"] / F) if F else 0, False)

    r = 6
    for m in order:
        sm.cell(r, 1, m)
        write_row(r, r, agg[m]); r += 1
    sm.cell(r, 1, "Grand Total")
    write_row(r, None, grand)
    for rr in range(6, r + 1):
        for cc in range(1, 10):
            cell = sm.cell(rr, cc); cell.border = BORDER
            if cc <= 7: cell.number_format = NUMACC
            else: cell.number_format = "0.0%"; cell.alignment = Alignment(horizontal="center")
    sm.cell(r, 1).alignment = Alignment(horizontal="left")
    return sm


def _run(wb, tol):
    """Mutate wb in place; return (stats, cache)."""
    ws, hrow, last, cmap, rows = _parse(wb, tol)
    planL = get_column_letter(cmap["plan"]); arrL = get_column_letter(cmap["arrdate"])
    poL = get_column_letter(cmap["poqty"]); arrqL = get_column_letter(cmap["arrqty"])
    cache = {}
    periodL, tolL, gapL, remL = _append_helpers(ws, hrow, rows, tol, planL, arrL, cache)
    agg = {}
    for x in rows:
        if not x["plan"]: continue
        m = MONTHS[x["plan"].month - 1]
        d = agg.setdefault(m, dict(lateL=0, lateS=0, ontL=0, ontS=0))
        if x["on_time"]: d["ontL"] += x["poqty"]; d["ontS"] += x["arrqty"]
        else: d["lateL"] += x["poqty"]; d["lateS"] += x["arrqty"]
    order = [m for m in MONTHS if m in agg]
    grand = {k: sum(agg[m][k] for m in order) for k in ("lateL","lateS","ontL","ontS")}
    title = str(next((x["pname"] for x in rows if x["pname"]), "PRINCIPAL")).strip()
    yrs = [x["plan"].year for x in rows if x["plan"]]
    year = max(set(yrs), key=yrs.count) if yrs else datetime.date.today().year
    if "SUMMARY" in wb.sheetnames: wb.remove(wb["SUMMARY"])
    _build_summary(wb, order, agg, grand, title, year, ws.title, poL, arrqL, periodL, remL, hrow + 1, last, cache)
    wb.active = wb.sheetnames.index("SUMMARY")
    try: wb.calculation.fullCalcOnLoad = True
    except Exception: pass
    F = grand["lateL"] + grand["ontL"]; G = grand["lateS"] + grand["ontS"]
    stats = dict(title=title, year=year, order=order, nlines=len(rows),
                 total_po=F, total_arr=G, fulfil=(G/F if F else 0), otif=(grand["ontS"]/F if F else 0),
                 late_po=grand["lateL"], late_arr=grand["lateS"], on_po=grand["ontL"], on_arr=grand["ontS"],
                 agg=agg, grand=grand)
    return stats, cache


def _fmt(v):
    if isinstance(v, float):
        return str(int(v)) if v.is_integer() else repr(v)
    return str(v)


def inject_cached_bytes(data, cache):
    """Return new xlsx bytes with <v> cached values inserted after <f>.

    Single pass over each sheet XML (O(size)) — replacing cells one by one with
    a fresh regex scan is O(cells x size) and crawls on large sheets."""
    if not cache: return data
    src = zipfile.ZipFile(io.BytesIO(data))
    names = src.namelist()
    wbx = src.read("xl/workbook.xml").decode("utf-8")
    rels = src.read("xl/_rels/workbook.xml.rels").decode("utf-8")
    relmap = {}
    for tag in re.findall(r"<Relationship\b[^>]*>", rels):
        mid = re.search(r'\bId="([^"]+)"', tag); mtg = re.search(r'\bTarget="([^"]+)"', tag)
        if mid and mtg: relmap[mid.group(1)] = mtg.group(1)
    sheetpath = {}
    for m in re.finditer(r'<sheet\b[^>]*name="([^"]+)"[^>]*r:id="([^"]+)"', wbx):
        tgt = relmap.get(m.group(2), "")
        if tgt: sheetpath[m.group(1)] = "xl/" + tgt.lstrip("/").replace("xl/", "", 1)
    data_map = {n: src.read(n) for n in names}
    src.close()

    cell_re = re.compile(r'<c r="([A-Z]+[0-9]+)"[^>]*>(.*?)</c>', re.S)
    empty_v_re = re.compile(r"<v[^>]*>.*?</v>", re.S)
    empty_v2_re = re.compile(r"<v\s*/>")
    for sheet, cells in cache.items():
        p = sheetpath.get(sheet)
        if not p or p not in data_map: continue
        xml = data_map[p].decode("utf-8")

        def repl(m, cells=cells):
            ref = m.group(1)
            hit = cells.get(ref)
            if hit is None: return m.group(0)
            inner = m.group(2)
            if "<f>" not in inner and "<f " not in inner: return m.group(0)
            val, is_str = hit
            inner2 = empty_v_re.sub("", inner)
            inner2 = empty_v2_re.sub("", inner2)
            tag = m.group(0)
            open_tag = tag[:tag.index(">") + 1]
            if is_str:
                open_tag = re.sub(r' t="[^"]*"', ' t="str"', open_tag) if ' t="' in open_tag else open_tag[:-1] + ' t="str">'
                v = "<v>%s</v>" % escape(str(val))
            else:
                v = "<v>%s</v>" % _fmt(val)
            return open_tag + inner2 + v + "</c>"

        data_map[p] = cell_re.sub(repl, xml).encode("utf-8")

    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for n in names: z.writestr(n, data_map[n])
    return out.getvalue()


def out_name(path):
    b = re.sub(r"\.xlsx?$", "", str(path).split("/")[-1].split("\\")[-1], flags=re.I)
    b = re.sub(r"[^0-9A-Za-z ]+", " ", b).strip()
    return f"2026_{b.upper()}_SL & OTIF.xlsx"


def process_bytes(data, tol=TOL, filename="upload.xlsx"):
    """Process an xlsx byte payload -> (out_bytes, stats, out_filename)."""
    wb = load_workbook(io.BytesIO(data))  # keeps formatting
    stats, cache = _run(wb, tol)
    buf = io.BytesIO(); wb.save(buf)
    return inject_cached_bytes(buf.getvalue(), cache), stats, out_name(filename)


def process_file(src, dst=None, tol=TOL):
    with open(src, "rb") as fh: data = fh.read()
    out, stats, name = process_bytes(data, tol, src)
    if dst is None:
        dst = name
    with open(dst, "wb") as fh: fh.write(out)
    return dst, stats, name
